import json, numpy as np, matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
from openpyxl import load_workbook
wb=load_workbook("out/MMForgings_Credit_Rating_Analysis.xlsx",data_only=True); an=wb["Analysis"]; sd=wb["Source_Data"]
lab={str(an[f"A{r}"].value):r for r in range(1,an.max_row+1) if an[f"A{r}"].value}
g=lambda l,cols="BCDEFG":[an[f"{c}{lab[l]}"].value for c in cols]
Y=["FY22","FY23","FY24","FY25","FY26","FY27E"]; N,GD,GR,RD="#1F3864","#C9A227","#7F7F7F","#C00000"
def st(ax,t):
    ax.set_title(t,loc="left",fontsize=11,fontweight="bold",color=N); ax.spines[["top","right"]].set_visible(False); ax.grid(axis="y",alpha=.3); ax.tick_params(labelsize=8)
x=np.arange(6); fig,ax=plt.subplots(figsize=(7,3.2))
ax.bar(x-0.2,g("Gross debt (closing)"),0.4,color=N,label="Gross debt"); ax.bar(x+0.2,g("PBILDT (operating profit)"),0.4,color=GD,label="PBILDT")
ax2=ax.twinx(); ax2.plot(x,g("Total debt / PBILDT"),color=RD,marker="o",label="Debt/PBILDT (rhs)"); ax2.axhline(3,ls="--",color=RD,lw=.8); ax2.text(5.4,3.05,"CARE upgrade <3.0x",fontsize=7,color=RD,ha="right")
ax2.spines[["top"]].set_visible(False); ax2.tick_params(labelsize=8); ax.set_xticks(x,Y); ax.legend(fontsize=7,frameon=False,loc="upper left"); ax2.legend(fontsize=7,frameon=False,loc="upper center")
st(ax,"Debt grew ₹479 Cr over FY23–26; PBILDT grew ₹7 Cr (₹ Cr)"); fig.tight_layout(); fig.savefig("out/charts/1_debt_pbildt.png",dpi=200); plt.close()
fig,ax=plt.subplots(figsize=(7,3.0)); ax.plot(x,g("Interest coverage (PBILDT / interest)"),marker="o",color=N,label="Interest coverage (x)")
ax2=ax.twinx(); r=g("ROCE (PBIT / avg capital employed)"); ax2.plot(x[1:],[v*100 for v in r[1:]],marker="s",color=GD,label="ROCE % (rhs)")
ax2.spines[["top"]].set_visible(False); ax2.tick_params(labelsize=8); ax.set_xticks(x,Y); ax.legend(fontsize=7,frameon=False,loc="upper right"); ax2.legend(fontsize=7,frameon=False,loc="lower left")
st(ax,"Coverage and returns compressed during the capex cycle"); fig.tight_layout(); fig.savefig("out/charts/2_coverage.png",dpi=200); plt.close()
q=[sd.cell(sd_r:=None or 0,1) for _ in []]
L={str(sd[f"A{r}"].value):r for r in range(1,sd.max_row+1) if sd[f"A{r}"].value}
qh=[sd.cell(L["Quarter"],c).value for c in range(2,15)]; qs=[sd.cell(L["Sales"],c).value for c in range(2,15)]; qm=[sd.cell(L["OPM %"],c).value for c in range(2,15)]
fig,ax=plt.subplots(figsize=(7,3.0)); xi=np.arange(13); ax.bar(xi,qs,color=[N]*12+[GD]); ax.set_ylim(300,450)
ax2=ax.twinx(); ax2.plot(xi,[m*100 for m in qm],color=RD,marker="o"); ax2.set_ylabel("OPM %",fontsize=8); ax2.spines[["top"]].set_visible(False); ax2.tick_params(labelsize=8)
ax.set_xticks(xi,qh,rotation=45,fontsize=7); st(ax,"Quarterly sales (₹ Cr, bars) and OPM % (line)")
fig.tight_layout(); fig.savefig("out/charts/3_quarterly.png",dpi=200); plt.close()
b=json.load(open("out/bridge.json")); fig,ax=plt.subplots(figsize=(7,3.0))
labels=["Model","Net liquid\ninvestments","CARE debt\ndefinition","FY27E only","CARE\naccruals"]; sc=[s[1] for s in b]
bars=ax.bar(labels,sc,color=[N,N,N,N,GD]); ax.set_ylim(6,7.8); ax.axhline(7.5,ls="--",color=RD,lw=1); ax.text(4.45,7.53,"CARE A threshold (7.5)",fontsize=7,color=RD,ha="right")
for bb,s in zip(bars,b): ax.text(bb.get_x()+bb.get_width()/2,s[1]+0.03,f"{s[1]:.2f}\n{s[2]}",ha="center",fontsize=8)
ax.tick_params(axis="x",labelsize=7); st(ax,"Notch bridge: model score after each cumulative adjustment"); fig.tight_layout(); fig.savefig("out/charts/4_bridge.png",dpi=200); plt.close()
