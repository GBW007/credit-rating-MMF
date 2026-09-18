from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
import json
F="Arial"; NAVY="1F3864"
BLUE=Font(name=F,size=10,color="0000FF"); BLK=Font(name=F,size=10); GRN=Font(name=F,size=10,color="008000")
BOLD=Font(name=F,size=10,bold=True); HDR=Font(name=F,size=10,bold=True,color="FFFFFF"); SUB=Font(name=F,size=9,italic=True,color="595959")
FH=PatternFill("solid",fgColor=NAVY); FS=PatternFill("solid",fgColor="D9E1F2"); FY=PatternFill("solid",fgColor="FFFF00")
NUM='#,##0;(#,##0);"-"'; NUM1='#,##0.0;(#,##0.0);"-"'; PCT='0.0%;(0.0%);"-"'; X='0.00"x";(0.00"x");"-"'; D0='0'
thin=Side(style="thin",color="BFBFBF")
wb=Workbook()
def sheet(n,t,s,wA=46):
    ws=wb.create_sheet(n); ws.sheet_view.showGridLines=False
    ws["A1"]=t; ws["A1"].font=Font(name=F,size=14,bold=True,color=NAVY); ws["A2"]=s; ws["A2"].font=SUB
    ws.column_dimensions["A"].width=wA
    for c in "BCDEFGHIJKLMNOP": ws.column_dimensions[c].width=11
    return ws
def put(ws,a,v,font=BLK,fmt=None,bold=False,fill=None):
    ws[a]=v; ws[a].font=Font(name=F,size=10,bold=True,color=font.color) if bold else font
    if fmt: ws[a].number_format=fmt
    if fill: ws[a].fill=fill
def band(ws,r,t,last="H"):
    for c in "ABCDEFGHIJKLMNOP"[: "ABCDEFGHIJKLMNOP".index(last)+1]: ws[f"{c}{r}"].fill=FH
    ws[f"A{r}"]=t; ws[f"A{r}"].font=HDR

YRS=["FY22","FY23","FY24","FY25","FY26"]; YC=list("BCDEF"); EST="G"  # G = FY27E
# ============================ SOURCE DATA
sd=sheet("Source_Data","M M Forgings Ltd (NSE: MMFL / BSE: 522241) — Source Data","Consolidated, ₹ crore, fiscal years ending 31 March. Blue = sourced input. Column H cites the source for each row.")
SD={}; r=4
put(sd,f"A{r}","₹ crore",BOLD)
for c,y in zip(YC,YRS): put(sd,f"{c}{r}",y,BOLD,fill=FS)
put(sd,f"H{r}","Source",BOLD,fill=FS); sd.column_dimensions["H"].width=60; sd[f"A{r}"].fill=FS
r+=1
SCR="Screener.in consolidated (data: C-MOTS), accessed 18-Sep-2026"
data=[("PROFIT & LOSS",None,None),
("sales","Sales / total operating income",[1140,1462,1563,1525,1590]),("exp","Operating expenses",[931,1189,1270,1229,1310]),
("oi","Other income",[16,12,22,23,16]),("int","Interest (finance cost)",[28,30,42,60,78]),("dep","Depreciation",[68,79,83,93,103]),
("pbt","Profit before tax",[129,177,189,166,115]),("pat","Net profit (PAT)",[91,128,135,123,99]),("payout","Dividend payout %",[0.16,0.11,0.14,0.16,0.20]),
("BALANCE SHEET (31 March)",None,None),
("eq","Equity share capital",[24,24,24,48,48]),("res","Reserves",[535,649,765,845,927]),("debt","Borrowings (gross)",[685,756,931,1185,1235]),
("othl","Other liabilities",[170,232,249,271,286]),("fa","Net fixed assets",[708,804,914,964,1101]),("cwip","Capital work-in-progress",[36,63,124,383,335]),
("inv","Investments (non-current, per Screener)",[22,22,23,4,10]),("otha","Other assets (current assets etc.)",[647,772,908,998,1051]),
("CASH FLOW",None,None),
("cfo","Cash from operating activity",[150,160,149,182,228]),("cfi","Cash from investing activity",[-126,-187,-232,-358,-182]),("cff","Cash from financing activity",[14,16,77,185,-38]),
("WORKING-CAPITAL DAYS",None,None),
("dd","Debtor days",[50,49,64,85,96]),("id","Inventory days",[147,163,176,188,171]),("pd","Days payable",[79,89,82,95,103])]
for k,l,v in data:
    if v is None: band(sd,r,k); r+=1; continue
    put(sd,f"A{r}",l)
    for c,x in zip(YC,v): put(sd,f"{c}{r}",x,BLUE,PCT if k=="payout" else NUM)
    put(sd,f"H{r}",SCR,SUB); SD[k]=r; r+=1
r+=1; band(sd,r,"CARE RATINGS — PUBLISHED FIGURES & DISCLOSURES (press release 30-Jun-2026)"); r+=1
CARE="CARE Ratings press release, M M Forgings Ltd, 30-Jun-2026"
care=[("c_toi","Total operating income (CARE)",[None,None,1563.07,1525.18,1589.87],NUM1),("c_pbildt","PBILDT (CARE)",[None,None,292.83,296.08,279.98],NUM1),
("c_pat","PAT (CARE)",[None,None,135.04,121.86,98.67],NUM1),("c_gear","Overall gearing (CARE)",[None,None,1.18,1.33,1.10],X),
("c_icr","Interest coverage (CARE table)",[None,None,6.92,4.95,3.58],X),("c_liq","Liquid investments + cash & bank",[None,None,None,218,225],NUM),
("c_dpb","Total debt / PBILDT (CARE text)",[None,None,None,4.00,3.84],X),("c_roce","ROCE (CARE text)",[None,None,None,0.1162,0.0916],PCT)]
for k,l,v,fmt in care:
    put(sd,f"A{r}",l)
    for c,x in zip(YC,v):
        if x is not None: put(sd,f"{c}{r}",x,BLUE,fmt)
    put(sd,f"H{r}",CARE,SUB); SD[k]=r; r+=1
single=[("rep_lo","Term-debt repayment p.a., next 3 yrs — low (₹ Cr)",150),("rep_hi","Term-debt repayment p.a., next 3 yrs — high (₹ Cr)",170),
("acc_lo","CARE-envisaged cash accruals p.a. — low",230),("acc_hi","CARE-envisaged cash accruals p.a. — high",350),
("wcu","Fund-based WC limit utilisation (12m to May-26)",0.80),("top10","Top-10 customer share (9MFY26)",0.62),("cv","Commercial-vehicle share of sales (FY25)",0.76),
("exp_sh","Export share of sales (FY26: US 10%, Europe 20%, others)",0.35),("capex3","Capex last 3 years (₹ Cr)",850),("capex3d","of which debt-funded (₹ Cr)",610),
("cap","Forging capacity, Dec-2025 (MT p.a.)",130000),("up_gear","CARE upgrade trigger: gearing below",0.80),("up_dpb","CARE upgrade trigger: debt/PBILDT below",3.0),
("up_roce","CARE upgrade trigger: ROCE above",0.15),("dn_gear","CARE downgrade trigger: gross gearing above",1.5),("icr_txt","Interest coverage FY26 quoted in CARE text",3.98)]
for k,l,v in single:
    put(sd,f"A{r}",l); fmt=PCT if isinstance(v,float) and v<1 else (X if k in("up_dpb","icr_txt") else NUM)
    if k in("up_gear","dn_gear"): fmt=X
    put(sd,f"F{r}",v,BLUE,fmt); put(sd,f"H{r}",CARE,SUB); SD[k]=r; r+=1
r+=1; band(sd,r,"MARKET DATA (Screener.in, 18-Sep-2026)"); r+=1
for k,l,v,fmt in [("mcap","Market capitalisation (₹ Cr)",3042,NUM),("prom","Promoter holding",0.5633,PCT),("pledge","Promoter shares pledged",0,PCT)]:
    put(sd,f"A{r}",l); put(sd,f"F{r}",v,BLUE,fmt); put(sd,f"H{r}",SCR+("; no pledge reported" if k=="pledge" else ""),SUB); SD[k]=r; r+=1
r+=1; band(sd,r,"QUARTERLY TREND (consolidated, ₹ crore)",last="P"); r+=1
Q=["Jun-23","Sep-23","Dec-23","Mar-24","Jun-24","Sep-24","Dec-24","Mar-25","Jun-25","Sep-25","Dec-25","Mar-26","Jun-26"]
qd={"q_sales":[370,397,399,398,382,398,374,371,362,385,414,430,420],"q_op":[65,75,75,78,73,77,73,73,63,67,69,81,75],
"q_oi":[5,5,6,6,6,7,4,6,8,4,3,1,63],"q_int":[10,9,11,13,15,16,15,14,18,21,20,18,16],"q_pat":[29,35,34,37,30,32,27,33,19,17,18,45,90]}
put(sd,f"A{r}","Quarter",BOLD,fill=FS)
QC=[chr(ord("B")+i) for i in range(13)]
for c,q in zip(QC,Q): put(sd,f"{c}{r}",q,BOLD,fill=FS)
SD["q_hdr"]=r; r+=1
for k,l in [("q_sales","Sales"),("q_op","Operating profit (PBILDT)"),("q_oi","Other income"),("q_int","Interest"),("q_pat","PAT")]:
    put(sd,f"A{r}",l)
    for c,x in zip(QC,qd[k]): put(sd,f"{c}{r}",x,BLUE,NUM)
    SD[k]=r; r+=1
put(sd,f"A{r}","OPM %")
for c in QC: put(sd,f"{c}{r}",f"={c}{SD['q_op']}/{c}{SD['q_sales']}",BLK,PCT)
SD["q_opm"]=r; r+=1
put(sd,f"A{r}","Note: Jun-26 other income of ₹63 Cr is ~10x the quarterly run-rate and tax was ~1%; treated as non-recurring in this analysis pending disclosure.",SUB)
sd[f"B{SD['q_oi']+0}"].comment=None
sd[f"N{SD['q_oi']}"].fill=FY
sd.freeze_panes="B5"
def S(k,c): return f"Source_Data!{c}{SD[k]}"
def S1(k): return f"Source_Data!$F${SD[k]}"

# ============================ ASSUMPTIONS FY27E
wa=sheet("FY27E_Base","FY27E Base Case (analyst estimate)","Forward view used alongside FY26 actuals, as rating agencies rate on expected performance. Every assumption is blue and justified.")
A={}; r=4
ass=[("g","Revenue growth FY27E",0.08,PCT,"Q1FY27 sales +16% YoY (₹420 Cr vs ₹362 Cr); haircut to 8% for CV cyclicality and US tariff uncertainty"),
("m","PBILDT margin FY27E",0.18,PCT,"Q4FY26 18.9%, Q1FY27 17.9%; CARE: margin recovery expected as new presses ramp up"),
("oi","Other income (recurring) ₹ Cr",16,NUM,"FY26 level; excludes the ₹63 Cr Q1FY27 one-off"),
("dep","Depreciation ₹ Cr",118,NUM,"Q1FY27 run-rate ₹30 Cr/qtr less ramp"),
("rate","Average cost of debt",0.065,PCT,"FY26 interest ₹78 Cr / avg debt ₹1,210 Cr = 6.4%"),
("rep","Term-debt repayment FY27E ₹ Cr",160,NUM,"Mid-point of CARE-disclosed ₹150–170 Cr"),
("newdebt","New debt drawn FY27E ₹ Cr",100,NUM,"Part-funding of ₹230 Cr 16,500T press (company); CARE: new debt capex aligned to repayments"),
("tax","Tax rate",0.25,PCT,"Normalised; FY26 14% was below normal"),
("div","Dividend ₹ Cr",20,NUM,"FY26 payout ~20% of PAT")]
put(wa,"A3","Assumption",BOLD); put(wa,"B3","Value",BOLD); put(wa,"C3","Rationale",BOLD); wa.column_dimensions["C"].width=95
for k,l,v,fmt,n in ass:
    put(wa,f"A{r}",l); put(wa,f"B{r}",v,BLUE,fmt,fill=FY if k in("g","m") else None); put(wa,f"C{r}",n,SUB); A[k]=r; r+=1
def AS(k): return f"FY27E_Base!$B${A[k]}"

# ============================ ANALYSIS
an=sheet("Analysis","Financial Statement & Ratio Analysis","FY22–FY26 actual (linked, green) + FY27E base case. Ratios use CARE definitions where stated (PBILDT = operating profit before other income).")
AN={}; r=4
put(an,f"A{r}","₹ crore",BOLD,fill=FS)
for c,y in zip(YC+[EST],YRS+["FY27E"]): put(an,f"{c}{r}",y,BOLD,fill=FS)
r+=1
def row(k,l,hist,est=None,fmt=NUM,bold=False,cols=YC):
    global r
    AN[k]=r; put(an,f"A{r}",l,bold=bold)
    for c in cols:
        f=hist(c); put(an,f"{c}{r}",f,GRN if (isinstance(f,str) and f.startswith("=Source")) else BLK,fmt,bold=bold)
    if est is not None: put(an,f"{EST}{r}",est,BLK,fmt,bold=bold)
    r+=1
L=lambda k:(lambda c:f"={S(k,c)}")
def p(c): return chr(ord(c)-1)
band(an,r,"PROFIT & LOSS AND CASH ACCRUALS",last="G"); r+=1
row("sales","Total operating income",L("sales"),f"=F{r}*(1+{AS('g')})",bold=True)
row("pbildt","PBILDT (operating profit)",lambda c:f"={S('sales',c)}-{S('exp',c)}",f"=G{AN['sales']}*{AS('m')}",bold=True)
row("m","PBILDT margin",lambda c:f"={c}{AN['pbildt']}/{c}{AN['sales']}",f"=G{AN['pbildt']}/G{AN['sales']}",PCT)
row("oi","Other income",L("oi"),f"={AS('oi')}")
row("int","Interest",L("int"),f"=AVERAGE(F{r+6},G{r+6})*{AS('rate')}")
row("dep","Depreciation",L("dep"),f"={AS('dep')}")
row("pbt","PBT",lambda c:f"={c}{AN['pbildt']}+{c}{AN['oi']}-{c}{AN['int']}-{c}{AN['dep']}",f"=G{AN['pbildt']}+G{AN['oi']}-G{AN['int']}-G{AN['dep']}")
row("pat","PAT",L("pat"),f"=G{AN['pbt']}*(1-{AS('tax')})",bold=True)
row("acc","Gross cash accruals (PAT + depreciation)",lambda c:f"={c}{AN['pat']}+{c}{AN['dep']}",f"=G{AN['pat']}+G{AN['dep']}",bold=True)
row("debt","Gross debt (closing)",L("debt"),f"=F{r}-{AS('rep')}+{AS('newdebt')}")
row("tnw","Net worth (equity + reserves)",lambda c:f"={S('eq',c)}+{S('res',c)}",f"=F{r}+G{AN['pat']}-{AS('div')}")
row("liq","Liquid investments + cash (CARE)",lambda c:(f"={S('c_liq',c)}" if c in "EF" else "=0"),f"=F{r}")
row("ndebt","Net debt",lambda c:f"={c}{AN['debt']}-{c}{AN['liq']}",f"=G{AN['debt']}-G{AN['liq']}")
# fix interest avg ref
an[f"G{AN['int']}"]=f"=AVERAGE(F{AN['debt']},G{AN['debt']})*{AS('rate')}"
row("ce","Capital employed (net worth + debt)",lambda c:f"={c}{AN['tnw']}+{c}{AN['debt']}",f"=G{AN['tnw']}+G{AN['debt']}")
row("cfo","Cash from operations",L("cfo"),None)
row("capex","Capex (CFI, proxy)",lambda c:f"=-{S('cfi',c)}",None)
row("fcf","Free cash flow (CFO − capex)",lambda c:f"={c}{AN['cfo']}-{c}{AN['capex']}",None,bold=True)
r+=1; band(an,r,"KEY CREDIT RATIOS",last="G"); r+=1
row("g","Revenue growth",lambda c:(f"={c}{AN['sales']}/{p(c)}{AN['sales']}-1"),f"=G{AN['sales']}/F{AN['sales']}-1",PCT,cols=YC[1:])
row("gear","Overall gearing (gross debt / net worth)",lambda c:f"={c}{AN['debt']}/{c}{AN['tnw']}",f"=G{AN['debt']}/G{AN['tnw']}",X)
row("ngear","Net gearing",lambda c:f"={c}{AN['ndebt']}/{c}{AN['tnw']}",f"=G{AN['ndebt']}/G{AN['tnw']}",X)
row("dpb","Total debt / PBILDT",lambda c:f"={c}{AN['debt']}/{c}{AN['pbildt']}",f"=G{AN['debt']}/G{AN['pbildt']}",X)
row("ndpb","Net debt / PBILDT",lambda c:f"={c}{AN['ndebt']}/{c}{AN['pbildt']}",f"=G{AN['ndebt']}/G{AN['pbildt']}",X)
row("icr","Interest coverage (PBILDT / interest)",lambda c:f"={c}{AN['pbildt']}/{c}{AN['int']}",f"=G{AN['pbildt']}/G{AN['int']}",X)
row("roce","ROCE (PBIT / avg capital employed)",lambda c:(f"=({c}{AN['pbt']}+{c}{AN['int']})/AVERAGE({c}{AN['ce']},{p(c)}{AN['ce']})"),f"=(G{AN['pbt']}+G{AN['int']})/AVERAGE(G{AN['ce']},F{AN['ce']})",PCT,cols=YC[1:])
row("roe","ROE (PAT / avg net worth)",lambda c:f"={c}{AN['pat']}/AVERAGE({c}{AN['tnw']},{p(c)}{AN['tnw']})",f"=G{AN['pat']}/AVERAGE(G{AN['tnw']},F{AN['tnw']})",PCT,cols=YC[1:])
row("cfoeb","CFO / PBILDT (cash conversion)",lambda c:f"={c}{AN['cfo']}/{c}{AN['pbildt']}",None,PCT)
row("dcfo","Total debt / CFO (years)",lambda c:f"={c}{AN['debt']}/{c}{AN['cfo']}",None,X)
row("ccc","Cash conversion cycle (days)",lambda c:f"={S('dd',c)}+{S('id',c)}-{S('pd',c)}",None,D0)
row("fat","Asset turnover (sales / (NFA + CWIP))",lambda c:f"={c}{AN['sales']}/({S('fa',c)}+{S('cwip',c)})",None,X)
row("cwipsh","CWIP as % of gross block proxy (NFA + CWIP)",lambda c:f"={S('cwip',c)}/({S('fa',c)}+{S('cwip',c)})",None,PCT)
r+=1; put(an,f"A{r}","Summary metrics",BOLD); r+=1
summ=[("cagr3","Revenue CAGR FY23–FY26",f"=(F{AN['sales']}/C{AN['sales']})^(1/3)-1",PCT),
("pbildt_cagr","PBILDT CAGR FY23–FY26",f"=(F{AN['pbildt']}/C{AN['pbildt']})^(1/3)-1",PCT),
("debt_add","Debt added FY23–FY26 (₹ Cr)",f"=F{AN['debt']}-C{AN['debt']}",NUM),
("pbildt_add","PBILDT added FY23–FY26 (₹ Cr)",f"=F{AN['pbildt']}-C{AN['pbildt']}",NUM),
("cumfcf","Cumulative FCF FY23–FY26 (₹ Cr)",f"=SUM(C{AN['fcf']}:F{AN['fcf']})",NUM)]
for k,l,f,fmt in summ: put(an,f"A{r}",l); put(an,f"C{r}",f,BLK,fmt,bold=True); AN[k]=r; r+=1
an.freeze_panes="B5"
def AA(k,c="F"): return f"Analysis!${c}${AN[k]}"

# ============================ DEBT SERVICE
ds=sheet("Debt_Service","Debt-Service Capacity — FY27E","Can cash accruals meet ₹150–170 Cr of annual term-debt repayment? Tested with and without the liquid-investment cushion.")
DS={}; r=4
hd=["Scenario","PBILDT","Interest","Tax","Cash accruals","Repayment","DSCR","Accruals / repayment","Liquid cover*"]
for i,h in enumerate(hd): put(ds,f"{chr(65+i)}{r}",h,BOLD,fill=FS)
ds.column_dimensions["A"].width=44
r+=1; sc=[("FY26 actual, mid repayment",f"={AA('pbildt')}",f"={AA('int')}",1.0,f"={S1('rep_lo')}/2+{S1('rep_hi')}/2"),
("FY27E base case",f"={AA('pbildt','G')}",f"={AA('int','G')}",1.0,f"={AS('rep')}"),
("FY27E: PBILDT −15% (CV downturn)",f"={AA('pbildt','G')}*0.85",f"={AA('int','G')}",1.0,f"={AS('rep')}"),
("FY27E: PBILDT −25%, repayment at high end",f"={AA('pbildt','G')}*0.75",f"={AA('int','G')}",1.0,f"={S1('rep_hi')}"),
("FY27E: PBILDT −25%, rates +200 bps, high repayment",f"={AA('pbildt','G')}*0.75",f"={AA('int','G')}+0.02*{AA('debt','G')}",1.0,f"={S1('rep_hi')}")]
DS0=r
for n,pb,it,_,rp in sc:
    put(ds,f"A{r}",n); put(ds,f"B{r}",pb,GRN,NUM); put(ds,f"C{r}",it,GRN,NUM)
    put(ds,f"D{r}",f"=MAX(0,B{r}+{AS('oi')}-C{r}-{AS('dep')})*{AS('tax')}",BLK,NUM)
    put(ds,f"E{r}",f"=B{r}+{AS('oi')}-C{r}-D{r}",BLK,NUM); put(ds,f"F{r}",rp,GRN,NUM)
    put(ds,f"G{r}",f"=(E{r}+C{r})/(C{r}+F{r})",BLK,X,bold=True); put(ds,f"H{r}",f"=E{r}/F{r}",BLK,X)
    put(ds,f"I{r}",f"=(E{r}+{AA('liq')})/F{r}",BLK,X)
    ds.conditional_formatting.add(f"G{r}",CellIsRule(operator="lessThan",formula=["1"],font=Font(color="C00000",bold=True)))
    r+=1
DS["base"]=DS0+1; DS["worst"]=r-1
r+=1
put(ds,f"A{r}","* Liquid cover = (cash accruals + ₹225 Cr liquid investments & cash) / repayment. DSCR here = (cash accruals + interest) / (interest + repayment), using all interest incl. working-capital interest — conservative.",SUB); r+=2
put(ds,f"A{r}","PBILDT decline that takes accruals/repayment to 1.0x (FY27E, pre-tax approx.)",BOLD)
put(ds,f"G{r}",f"=1-(G{DS['base']}-(E{DS['base']}-F{DS['base']})/(1-{AS('tax')}))/G{DS['base']}".replace(f"G{DS['base']}",f"B{DS['base']}",1),BLK,PCT,bold=True)
ds[f"G{r}"]=f"=((E{DS['base']}-F{DS['base']})/(1-{AS('tax')}))/B{DS['base']}"; ds[f"G{r}"].number_format=PCT; ds[f"G{r}"].font=BOLD
DS["cush"]=r

# ============================ RATING
rg=sheet("Rating","Credit Rating Scorecard — M M Forgings","Same hybrid methodology as the SVF case, recalibrated for a listed mid-corporate. Financial metrics blend FY26 actual (60%) and FY27E base (40%).",wA=40)
rg.column_dimensions["B"].width=8; rg.column_dimensions["C"].width=10; rg.column_dimensions["D"].width=10; rg.column_dimensions["E"].width=7; rg.column_dimensions["F"].width=8; rg.column_dimensions["G"].width=70
G={}; r=3
put(rg,"A3","Debt basis for leverage metrics (Gross / Net)",BOLD); put(rg,"C3","Gross",BLUE,fill=FY); G["basis"]=3
put(rg,"G3","Gross = conservative bank view; Net credits ₹225 Cr liquid investments (CARE's liquidity view)",SUB)
put(rg,"A4","Weight on FY26 actual (rest on FY27E)",BOLD); put(rg,"C4",0.6,BLUE,PCT,fill=FY); G["w"]=4
r=6
for c,h in zip("ABCDEFG",["Parameter","Weight","FY26","FY27E","Blend","Score","Rationale"]): put(rg,f"{c}{r}",h,BOLD,fill=FS)
r+=1
grids={"dpb":([0,1,1.75,2.5,3.0,3.5,4.25,5.5],[10,9,8,7,6,5,3,1]),"gear":([0,0.25,0.5,0.8,1.0,1.25,1.5,2.5],[10,9,8,7,6,5,4,1]),
"icr":([0,1.5,2,2.5,3,3.5,4.5,6],[1,3,5,6,7,8,9,10]),"acc":([0,0.8,1,1.2,1.4,1.7,2,2.5],[1,2,4,5,6,7,8,10]),
"roce":([0,0.05,0.08,0.10,0.12,0.15,0.18,0.22],[1,3,4,5,6,7,8,10]),"m":([0,0.05,0.08,0.10,0.12,0.15,0.18,0.22],[1,3,5,6,7,8,9,10]),
"scale":([0,100,250,500,1000,2500,5000,10000],[1,3,4,5,7,8,9,10]),"cfo":([0,0.3,0.5,0.6,0.7,0.8,0.9,1.0],[1,3,5,6,7,8,9,10])}
gr=40; GR={}
put(rg,f"I{gr-1}","Threshold grids (breakpoint → score)",BOLD)
for k,(bp,s) in grids.items():
    put(rg,f"I{gr}",k,BOLD)
    for j,(b,x) in enumerate(zip(bp,s)):
        c=chr(ord("J")+j); put(rg,f"{c}{gr}",b,BLUE,'0.00'); put(rg,f"{c}{gr+1}",x,BLUE,'0')
    GR[k]=gr; gr+=3
for c in "JKLMNOPQ": rg.column_dimensions[c].width=7
def sf(v,k): g=GR[k]; return f"=INDEX($J${g+1}:$Q${g+1},MATCH({v},$J${g}:$Q${g},1))"
band(rg,r,"A. FINANCIAL RISK (50%)",last="G"); r+=1
isnet=f'$C${G["basis"]}="Net"'
q=[("Debt / PBILDT",0.10,f"=IF({isnet},{AA('ndpb')},{AA('dpb')})",f"=IF({isnet},{AA('ndpb','G')},{AA('dpb','G')})","dpb",X,"CARE upgrade trigger < 3.0x"),
("Overall gearing",0.08,f"=IF({isnet},{AA('ngear')},{AA('gear')})",f"=IF({isnet},{AA('ngear','G')},{AA('gear','G')})","gear",X,"CARE: upgrade < 0.8x; downgrade > 1.5x"),
("Interest coverage (PBILDT / interest)",0.07,f"={AA('icr')}",f"={AA('icr','G')}","icr",X,"Fell from 6.9x (FY24) as debt-funded capex raised interest"),
("Cash accruals / term-debt repayment",0.08,f"={AA('acc')}/({S1('rep_lo')}/2+{S1('rep_hi')}/2)",f"={AA('acc','G')}/{AS('rep')}","acc",X,"Repayment ₹150–170 Cr p.a. (CARE)"),
("ROCE",0.06,f"={AA('roce')}",f"={AA('roce','G')}","roce",PCT,"Depressed by ₹335 Cr CWIP not yet earning"),
("PBILDT margin",0.04,f"={AA('m')}",f"={AA('m','G')}","m",PCT,"Range-bound 17.6–19.4%"),
("Operating scale (revenue ₹ Cr)",0.04,f"={AA('sales')}",f"={AA('sales','G')}","scale",NUM,"Scale supports absorption of fixed costs"),
("Cash conversion (CFO / PBILDT)",0.03,f"={AA('cfoeb')}",f"={AA('cfoeb')}","cfo",PCT,"FY26 actual used for both columns")]
q0=r
for l,w,a,b,k,fmt,n in q:
    put(rg,f"A{r}",l); put(rg,f"B{r}",w,BLUE,PCT); put(rg,f"C{r}",a,GRN,fmt); put(rg,f"D{r}",b,GRN,fmt)
    put(rg,f"E{r}",f"=$C${G['w']}*C{r}+(1-$C${G['w']})*D{r}",BLK,fmt); put(rg,f"F{r}",sf(f"E{r}",k),BLK,'0'); put(rg,f"G{r}",n,SUB); r+=1
q1=r-1
def qual(title,items):
    global r
    band(rg,r,title,last="G"); r+=1
    for l,w,s,n in items:
        put(rg,f"A{r}",l); put(rg,f"B{r}",w,BLUE,PCT); put(rg,f"F{r}",s,BLUE,'0',fill=FY); put(rg,f"G{r}",n,SUB); r+=1
b0=r
qual("B. BUSINESS RISK (25%)",[("Track record & market position",0.05,9,"Forging since 1974; two-decade client relationships; 130,000 MTPA capacity"),
("Product & geographic diversity",0.05,8,"0.1–120 kg parts; ~35% exports (US 10%, Europe 20%)"),
("Customer concentration",0.05,5,"Top-10 = 62% of TOI (9MFY26), down from 69%"),
("End-market cyclicality",0.05,4,"~76% commercial vehicles — the most cyclical auto segment"),
("Raw-material & tariff exposure",0.05,6,"Pass-through clauses with lag; US tariff hit volumes in Q2–Q3 FY26")])
qual("C. MANAGEMENT & GOVERNANCE (10%)",[("Promoter experience & commitment",0.05,8,"Second-generation MD (IIT Madras), 25+ yrs; promoter holding 56.3%, no pledge"),
("Governance & disclosure",0.05,7,"Listed; independent directors; regular concalls. CARE text/table ICR mismatch noted")])
qual("D. LIQUIDITY & FINANCIAL FLEXIBILITY (15%)",[("Liquid investments & cash",0.05,8,"₹225 Cr (Mar-26) ≈ 1.4x one year's repayment"),
("Working-capital headroom",0.05,7,"Fund-based utilisation ~80% (12m to May-26)"),
("Access to capital",0.05,8,"Listed, ~₹3,000 Cr market cap; long banking relationships")])
last=r-1; r+=1
put(rg,f"A{r}","Total weight (must = 100%)",BOLD); put(rg,f"B{r}",f"=SUM(B{q0}:B{last})",BLK,PCT,bold=True); r+=1
put(rg,f"A{r}","Financial sub-score (/10)"); put(rg,f"F{r}",f"=SUMPRODUCT(B{q0}:B{q1},F{q0}:F{q1})/SUM(B{q0}:B{q1})",BLK,'0.00'); G["fin"]=r; r+=1
put(rg,f"A{r}","WEIGHTED SCORE (/10)",BOLD); put(rg,f"F{r}",f"=SUMPRODUCT(B{q0}:B{last},F{q0}:F{last})",BLK,'0.00',bold=True); G["score"]=r; r+=1
put(rg,f"A{r}","MODEL RATING",BOLD); G["grade"]=r; r+=1
put(rg,f"A{r}","Indicative 1-yr PD (illustrative)"); G["pd"]=r; r+=1
put(rg,f"A{r}","CARE rating (30-Jun-2026)"); put(rg,f"F{r}","A",BLUE,bold=True); put(rg,f"G{r}","CARE A; Stable / CARE A1 — reaffirmed",SUB); G["care"]=r; r+=1
put(rg,f"A{r}","Notch difference (+ = model more conservative)",BOLD); G["diff"]=r; r+=2
put(rg,f"A{r}","Master scale",BOLD); r+=1
for c,h in zip("ABCD",["Min score","Grade","Rank","PD (illus.)"]): put(rg,f"{c}{r}",h,BOLD,fill=FS)
r+=1; m0=r
scale=[(0,"D",13,1.0),(3.0,"C",12,0.25),(4.0,"B",11,0.10),(4.6,"BB-",10,0.05),(5.0,"BB",9,0.035),(5.4,"BB+",8,0.025),(5.8,"BBB-",7,0.015),
(6.2,"BBB",6,0.01),(6.6,"BBB+",5,0.006),(7.0,"A-",4,0.004),(7.5,"A",3,0.0025),(8.0,"A+",2,0.0015),(8.5,"AA",1,0.0008)]
for mn,gd,rk,pdv in scale: put(rg,f"A{r}",mn,BLUE,'0.0'); put(rg,f"B{r}",gd,BLUE); put(rg,f"C{r}",rk,BLUE,'0'); put(rg,f"D{r}",pdv,BLUE,'0.00%'); r+=1
m1=r-1
put(rg,f"A{r}","PDs illustrative, broadly consistent with long-run Indian CRA default studies by band; not calibrated.",SUB)
rg[f"F{G['grade']}"]=f"=INDEX($B${m0}:$B${m1},MATCH(F{G['score']},$A${m0}:$A${m1},1))"; rg[f"F{G['grade']}"].font=BOLD; rg[f"F{G['grade']}"].fill=FY
rg[f"F{G['pd']}"]=f"=INDEX($D${m0}:$D${m1},MATCH(F{G['score']},$A${m0}:$A${m1},1))"; rg[f"F{G['pd']}"].number_format='0.00%'
rg[f"F{G['diff']}"]=f"=INDEX($C${m0}:$C${m1},MATCH(F{G['grade']},$B${m0}:$B${m1},0))-INDEX($C${m0}:$C${m1},MATCH(F{G['care']},$B${m0}:$B${m1},0))"; rg[f"F{G['diff']}"].font=BOLD
json.dump({"q0":q0,"last":last},open("rg.json","w"))

# ============================ CARE BENCHMARK
cb=sheet("CARE_Benchmark","Benchmarking Against CARE Ratings","Reconciles our computed figures to CARE's published numbers and measures the distance to CARE's stated upgrade / downgrade triggers.")
CB={}; r=4
for c,h in zip("ABCDEF",["Metric (FY26)","Our figure","CARE figure","Difference","Status","Comment"]): put(cb,f"{c}{r}",h,BOLD,fill=FS)
cb.column_dimensions["F"].width=80; r+=1
rec=[("Total operating income",f"={AA('sales')}",f"={S('c_toi','F')}",NUM1,"Screener rounds to ₹ Cr"),
("PBILDT",f"={AA('pbildt')}",f"={S('c_pbildt','F')}",NUM1,""),("PAT",f"={AA('pat')}",f"={S('c_pat','F')}",NUM1,""),
("Interest coverage",f"={AA('icr')}",f"={S('c_icr','F')}",X,"CARE's own text quotes 3.98x vs 3.58x in its table — the table reconciles to PBILDT/interest"),
("Overall gearing",f"={AA('gear')}",f"={S('c_gear','F')}",X,"FY24–FY25 reconcile; FY26 gap implies CARE counts ~₹160 Cr less debt (likely classification, e.g. leases/acceptances). We keep the higher, conservative figure"),
("Total debt / PBILDT",f"={AA('dpb')}",f"={S('c_dpb','F')}",X,"Same debt-definition gap as gearing"),
("ROCE",f"={AA('roce')}",f"={S('c_roce','F')}",PCT,"Definitions differ slightly (avg vs closing capital employed)")]
for l,a,b,fmt,n in rec:
    put(cb,f"A{r}",l); put(cb,f"B{r}",a,GRN,fmt); put(cb,f"C{r}",b,GRN,fmt); put(cb,f"D{r}",f"=B{r}-C{r}",BLK,fmt)
    put(cb,f"E{r}",f'=IF(ABS(D{r})<=0.03*ABS(C{r}),"Reconciles","Investigate")',BLK); put(cb,f"F{r}",n,SUB)
    cb.conditional_formatting.add(f"E{r}",CellIsRule(operator="equal",formula=['"Investigate"'],font=Font(color="C65911",bold=True))); r+=1
put(cb,f"A{r}","Implied CARE debt (gearing × our net worth)"); put(cb,f"B{r}",f"={S('c_gear','F')}*{AA('tnw')}",BLK,NUM); CB["impl"]=r; r+=1
put(cb,f"A{r}","Implied CARE debt (debt/PBILDT × PBILDT)"); put(cb,f"B{r}",f"={S('c_dpb','F')}*{S('c_pbildt','F')}",BLK,NUM); r+=2
for c,h in zip("ABCDEF",["CARE trigger","Threshold","FY26","FY27E","Met?","Distance (FY27E)"]): put(cb,f"{c}{r}",h,BOLD,fill=FS)
r+=1; CB["t0"]=r
trig=[("Upgrade: gearing below",f"={S1('up_gear')}",f"={AA('gear')}",f"={AA('gear','G')}","<",X),
("Upgrade: debt / PBILDT below",f"={S1('up_dpb')}",f"={AA('dpb')}",f"={AA('dpb','G')}","<",X),
("Upgrade: ROCE above",f"={S1('up_roce')}",f"={AA('roce')}",f"={AA('roce','G')}",">",PCT),
("Downgrade: gross gearing above",f"={S1('dn_gear')}",f"={AA('gear')}",f"={AA('gear','G')}",">",X)]
for l,t,a,b,op,fmt in trig:
    put(cb,f"A{r}",l); put(cb,f"B{r}",t,GRN,fmt); put(cb,f"C{r}",a,GRN,fmt); put(cb,f"D{r}",b,GRN,fmt)
    put(cb,f"E{r}",f'=IF(D{r}{op}B{r},"Yes","No")',BLK); put(cb,f"F{r}",f"=D{r}-B{r}",BLK,fmt); r+=1
put(cb,f"A{r}","Debt reduction needed for debt/PBILDT < 3.0x at FY27E PBILDT (₹ Cr)",BOLD); put(cb,f"D{r}",f"=MAX(0,{AA('debt','G')}-{S1('up_dpb')}*{AA('pbildt','G')})",BLK,NUM,bold=True); CB["need"]=r

# ============================ DASHBOARD
d=wb["Sheet"]; d.title="Dashboard"; d.sheet_view.showGridLines=False
d["A1"]="M M Forgings Ltd — Independent Credit Rating Analysis"; d["A1"].font=Font(name=F,size=16,bold=True,color=NAVY)
d["A2"]="NSE: MMFL | BSE: 522241 | Consolidated | Actual public data (Screener.in, CARE Ratings) | Prepared by Gaurav Waghmode, 18-Sep-2026"; d["A2"].font=SUB
d.column_dimensions["A"].width=44; d.column_dimensions["B"].width=16; d.column_dimensions["C"].width=3; d.column_dimensions["D"].width=40; d.column_dimensions["E"].width=16
r=4
for c in "AB": d[f"{c}{r}"].fill=FH
for c in "DE": d[f"{c}{r}"].fill=FH
put(d,"A4","RATING OUTCOME",HDR); put(d,"D4","KEY METRICS (FY26 → FY27E)",HDR)
L_=[("Model rating",f"=Rating!F{G['grade']}",None),("CARE rating (Jun-2026)",'="CARE A; Stable"',None),("Notch difference (model vs CARE)",f"=Rating!F{G['diff']}",'+0;-0;0'),
("Weighted score / 10",f"=Rating!F{G['score']}",'0.00'),("Financial sub-score / 10",f"=Rating!F{G['fin']}",'0.00'),("Indicative PD",f"=Rating!F{G['pd']}",'0.00%'),
("Debt basis used",f"=Rating!C{G['basis']}",None),("FY27E base-case DSCR",f"=Debt_Service!G{DS['base']}",X),("Worst-scenario DSCR",f"=Debt_Service!G{DS['worst']}",X),
("PBILDT cushion before accruals < repayment",f"=Debt_Service!G{DS['cush']}",PCT)]
R_=[("Revenue (₹ Cr)",'sales',NUM),("PBILDT margin",'m',PCT),("Gross debt (₹ Cr)",'debt',NUM),("Gearing",'gear',X),("Debt / PBILDT",'dpb',X),("Net debt / PBILDT",'ndpb',X),("Interest coverage",'icr',X),("ROCE",'roce',PCT),("Cash accruals (₹ Cr)",'acc',NUM)]
for i,(l,f,fmt) in enumerate(L_): put(d,f"A{5+i}",l); put(d,f"B{5+i}",f,BLK,fmt,bold=True)
for i,(l,k,fmt) in enumerate(R_):
    put(d,f"D{5+i}",l)
    tf={"0":'"#,##0"',X:'"0.00x"',PCT:'"0.0%"'}.get(fmt,'"#,##0"')
    tf='"#,##0"' if fmt==NUM else ('"0.00""x"""' if fmt==X else '"0.0%"')
    d[f"E{5+i}"]=f"=TEXT({AA(k)},{tf})&\" → \"&TEXT({AA(k,'G')},{tf})"; d[f"E{5+i}"].font=BOLD; d[f"E{5+i}"].alignment=Alignment(horizontal="right")
d["B5"].fill=FY
# chart data
r=17; put(d,f"A{r}","Chart data (linked)",SUB); r+=1
put(d,f"A{r}","Year",BOLD)
for j,y in enumerate(YRS+["FY27E"]): put(d,f"{chr(66+j)}{r}",y,BOLD)
for i,(l,k) in enumerate([("Gross debt (₹ Cr)","debt"),("PBILDT (₹ Cr)","pbildt"),("Debt / PBILDT (x)","dpb"),("Gearing (x)","gear")]):
    put(d,f"A{r+1+i}",l)
    for j,c in enumerate(YC+["G"]): put(d,f"{chr(66+j)}{r+1+i}",f"={AA(k,c)}",GRN,'0.00')
for c in "FG": d.column_dimensions[c].width=10
bc=BarChart(); bc.title="Debt vs PBILDT (₹ Cr)"; bc.height=7; bc.width=15
bc.add_data(Reference(d,min_col=1,max_col=7,min_row=r+1,max_row=r+2),titles_from_data=True,from_rows=True); bc.set_categories(Reference(d,min_col=2,max_col=7,min_row=r))
d.add_chart(bc,"A24")
lc=LineChart(); lc.title="Leverage (x)"; lc.height=7; lc.width=15
lc.add_data(Reference(d,min_col=1,max_col=7,min_row=r+3,max_row=r+4),titles_from_data=True,from_rows=True); lc.set_categories(Reference(d,min_col=2,max_col=7,min_row=r))
d.add_chart(lc,"D24")
put(d,"A40","Sources: Screener.in consolidated financials (C-MOTS data) and CARE Ratings press release dated 30-Jun-2026. Independent analysis for educational purposes; not investment advice or an official rating.",SUB)
json.dump({"AN":AN,"DS":DS,"G":G,"CB":CB,"SD":SD},open("map.json","w"))
wb.save("out/MMForgings_Credit_Rating_Analysis.xlsx"); print("saved")
