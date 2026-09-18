"""Replicates the Rating sheet scorecard in Python and runs a notch bridge from the model rating to CARE's rating."""
import json, sys
from openpyxl import load_workbook
wbp=sys.argv[1] if len(sys.argv)>1 else "out/MMForgings_Credit_Rating_Analysis.xlsx"
wb=load_workbook(wbp,data_only=True); rg=wb["Rating"]; an=wb["Analysis"]
lab={str(an[f"A{r}"].value):r for r in range(1,an.max_row+1) if an[f"A{r}"].value}
v=lambda l,c:float(an[f"{c}{lab[l]}"].value)
grids={}
for r in range(40,70,3):
    k=rg[f"I{r}"].value
    if k: grids[k]=([rg.cell(r,c).value for c in range(10,18)],[rg.cell(r+1,c).value for c in range(10,18)])
def score(x,k):
    bp,s=grids[k]; out=s[0]
    for b,sc in zip(bp,s):
        if x>=b: out=sc
    return out
qual=sum(rg[f"B{r}"].value*rg[f"F{r}"].value for r in range(17,29) if isinstance(rg[f"B{r}"].value,(int,float)) and isinstance(rg[f"F{r}"].value,(int,float)))
scale=[(0,"D"),(3.0,"C"),(4.0,"B"),(4.6,"BB-"),(5.0,"BB"),(5.4,"BB+"),(5.8,"BBB-"),(6.2,"BBB"),(6.6,"BBB+"),(7.0,"A-"),(7.5,"A"),(8.0,"A+"),(8.5,"AA")]
grade=lambda s:[g for m,g in scale if s>=m][-1]
F,E="F","G"
base=dict(pb=(v("PBILDT (operating profit)",F),v("PBILDT (operating profit)",E)),debt=(v("Gross debt (closing)",F),v("Gross debt (closing)",E)),
 tnw=(v("Net worth (equity + reserves)",F),v("Net worth (equity + reserves)",E)),icr=(v("Interest coverage (PBILDT / interest)",F),v("Interest coverage (PBILDT / interest)",E)),
 acc=(v("Gross cash accruals (PAT + depreciation)",F)/160,v("Gross cash accruals (PAT + depreciation)",E)/160),roce=(v("ROCE (PBIT / avg capital employed)",F),v("ROCE (PBIT / avg capital employed)",E)),
 m=(v("PBILDT margin",F),v("PBILDT margin",E)),rev=(v("Total operating income",F),v("Total operating income",E)),cfo=(v("CFO / PBILDT (cash conversion)",F),)*2,
 liq=225,w=0.6,net=False,debt_adj=0)
def run(p):
    w=p["w"]; bl=lambda t:w*t[0]+(1-w)*t[1]
    d=tuple(x-p["debt_adj"]-(p["liq"] if p["net"] else 0) for x in p["debt"])
    f=[(0.10,score(bl((d[0]/p["pb"][0],d[1]/p["pb"][1])),"dpb")),(0.08,score(bl((d[0]/p["tnw"][0],d[1]/p["tnw"][1])),"gear")),
       (0.07,score(bl(p["icr"]),"icr")),(0.08,score(bl(p["acc"]),"acc")),(0.06,score(bl(p["roce"]),"roce")),(0.04,score(bl(p["m"]),"m")),
       (0.04,score(bl(p["rev"]),"scale")),(0.03,score(bl(p["cfo"]),"cfo"))]
    s=sum(a*b for a,b in f)+qual; return round(s,2),grade(s)
steps=[]; p=dict(base); steps.append(("Model as built (gross debt, 60% FY26 / 40% FY27E)",*run(p)))
p["net"]=True; steps.append(("+ Net off ₹225 Cr liquid investments",*run(p)))
p["debt_adj"]=160; steps.append(("+ Adopt CARE's narrower debt definition (~₹160 Cr lower)",*run(p)))
p["w"]=0.0; steps.append(("+ Rate purely on FY27E (forward-looking, as CRAs do)",*run(p)))
p["acc"]=(p["acc"][0],290/160); steps.append(("+ Use CARE's envisaged accruals (mid-point ₹290 Cr)",*run(p)))
print(f"qualitative contribution {qual:.2f}")
for s in steps: print(f"{s[1]:5.2f}  {s[2]:5}  {s[0]}")
json.dump(steps,open("out/bridge.json","w"),indent=1)
